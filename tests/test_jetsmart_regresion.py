"""Regresion de la liquidacion de JetSmart contra la armada a mano (julio 2026).

Correr con:
    python -m unittest tests.test_jetsmart_regresion -v

Necesita los archivos reales en tests/fixtures/jetsmart/ (no se versionan,
son datos del cliente -- ver .gitignore). Sin ellos, los tests se saltean.

Dos niveles, a proposito:

1. CALCULO (celda por celda, al centavo): se alimenta la liquidacion con
   las guias YA CONCILIADAS por Anita (hoja GUIAS de LIQ_ECS_07-2026) y se
   compara cada celda de valor de las hojas LIQUIDACION y CVLP del Excel
   generado contra la liquidacion a mano. Esto valida las formulas aisladas
   del cruce con Ariel, que esta fuera de alcance.

2. EXPORT CRUDO: se corre sobre el export real (hoja "BD 07-2026") y se
   verifica que la diferencia contra la liquidacion a mano sea EXACTAMENTE
   la de las guias que el cruce con Ariel agrego/saco -- es decir, que no
   hay ninguna otra diferencia escondida en el calculo.
"""

import io
import unittest
from pathlib import Path

import openpyxl
import pandas as pd

from src.config import (
    COL_JS_CLIENTE,
    COL_JS_CREACION,
    COL_JS_DESTINO,
    COL_JS_EMPRESA,
    COL_JS_ESTADO,
    COL_JS_GUIA,
    COL_JS_KGS,
    COL_JS_ORIGEN,
    COL_JS_PRIORIDAD,
)
from src.jetsmart_builder import (
    build_jetsmart_guias,
    build_jetsmart_resumen,
    detect_jetsmart_period,
    load_jetsmart_export,
    unconfirmed_tarifa_activity,
    write_jetsmart_liquidacion,
)

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "jetsmart"
LIQ_FINAL = FIXTURES / "LIQ_ECS_07-2026_00000005.xlsx"
EXPORT = FIXTURES / "LIQUIDACION_07-2026.xlsx"

# Datos manuales del periodo, tal como estan en la liquidacion a mano:
# TC en LIQUIDACION!H5; vuelos en MANI EZE (27 expo + 29 impo) y MANI MDZ
# (38 filas de 35 USD).
TC_JULIO_2026 = 1485
VUELOS_JULIO_2026 = {"EZE": 56, "MDZ": 38}
PERIODO = ("JUL", "26")

# Todas las celdas con valor de la hoja LIQUIDACION a mano (columna B:
# resumen; E5:I16: tabla por estacion) y de la hoja CVLP (columna C).
LIQUIDACION_CELDAS = ["B3", "B4", "B6", "B7", "B8", "B9", "B10", "B11", "B13"] + [
    f"{col}{row}" for row in range(5, 16) for col in "FGHI"
] + ["F16", "I16"]
LIQUIDACION_ETIQUETAS = [f"E{row}" for row in range(5, 17)]
CVLP_CELDAS = ["C2", "C3", "C4", "C9", "C10", "C11", "C12", "C14", "C15", "C16", "C19", "C21", "C23"]


def _sheet(wb: openpyxl.Workbook, name: str):
    # La planilla a mano tiene espacios al final de los nombres de hoja
    # ("LIQUIDACION ", "GUIAS "); la generada no.
    for ws in wb.worksheets:
        if ws.title.strip() == name:
            return ws
    raise KeyError(name)


def _as_number(value) -> float:
    return 0.0 if value is None else float(value)


def _guias_conciliadas_como_export() -> pd.DataFrame:
    """Hoja GUIAS de la liquidacion a mano, renombrada a las columnas del
    export (Nro guía -> # Guía, Kg -> KGs, Ingreso total flete -> $ Prioridad)."""
    g = pd.read_excel(LIQ_FINAL, sheet_name="GUIAS ", usecols=range(15))
    g = g[g["AWB"].notna()]
    return pd.DataFrame({
        COL_JS_CREACION: "15/07/2026",
        COL_JS_GUIA: g["Nro guía"],
        COL_JS_CLIENTE: g["Cliente"],
        COL_JS_ORIGEN: g["Origen"],
        COL_JS_DESTINO: g["Destino"],
        COL_JS_ESTADO: "Entregada",
        COL_JS_KGS: g["Kg"],
        COL_JS_PRIORIDAD: g["Ingreso total flete"],
        COL_JS_EMPRESA: "Jetsmart",
    }).reset_index(drop=True)


@unittest.skipUnless(LIQ_FINAL.exists() and EXPORT.exists(), "faltan los archivos de tests/fixtures/jetsmart/")
class TestCalculoCeldaPorCelda(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        guias = build_jetsmart_guias(_guias_conciliadas_como_export())
        cls.resumen = build_jetsmart_resumen(guias, TC_JULIO_2026, VUELOS_JULIO_2026)
        buffer = io.BytesIO()
        write_jetsmart_liquidacion(guias, cls.resumen, PERIODO, buffer)
        buffer.seek(0)
        cls.generado = openpyxl.load_workbook(buffer)
        cls.referencia = openpyxl.load_workbook(LIQ_FINAL, data_only=True)

    def _comparar(self, hoja: str, celdas: list[str]):
        gen, ref = _sheet(self.generado, hoja), _sheet(self.referencia, hoja)
        for celda in celdas:
            with self.subTest(hoja=hoja, celda=celda):
                # Tolerancia de medio centavo: el Excel generado guarda los
                # montos redondeados a 2 decimales; la planilla a mano, sin
                # redondear.
                self.assertAlmostEqual(_as_number(gen[celda].value), _as_number(ref[celda].value), delta=0.005)

    def test_hoja_liquidacion(self):
        self._comparar("LIQUIDACION", LIQUIDACION_CELDAS)

    def test_hoja_liquidacion_estaciones_en_el_mismo_orden(self):
        gen, ref = _sheet(self.generado, "LIQUIDACION"), _sheet(self.referencia, "LIQUIDACION")
        for celda in LIQUIDACION_ETIQUETAS:
            with self.subTest(celda=celda):
                self.assertEqual(gen[celda].value, ref[celda].value)

    def test_hoja_cvlp(self):
        self._comparar("CVLP", CVLP_CELDAS)

    def test_resumen_sin_redondear(self):
        # El dict (que es lo que se persiste y se muestra en la app) tiene
        # que coincidir tambien, no solo el Excel.
        ref = _sheet(self.referencia, "LIQUIDACION")
        self.assertAlmostEqual(self.resumen["ventas_totales"], ref["B3"].value, places=4)
        self.assertAlmostEqual(self.resumen["total_wcs"], ref["B13"].value, places=4)
        self.assertAlmostEqual(self.resumen["cvlp"]["total_final"], _sheet(self.referencia, "CVLP")["C23"].value, places=4)

    def test_tarifa_confirmada_sin_aviso_en_julio(self):
        # Todas las estaciones con kilos en julio 2026 tienen la tarifa
        # confirmada (ver JETSMART_TARIFA_CONFIRMADA_STATIONS).
        self.assertEqual(unconfirmed_tarifa_activity(self.resumen), {})

    def test_tarifa_no_confirmada_se_avisa_sin_bloquear(self):
        # Una estacion que no esta en config (todas las conocidas estan
        # confirmadas): si tiene kilos, se calcula igual con la tarifa
        # default y se avisa.
        export = _guias_conciliadas_como_export().head(1).assign(**{COL_JS_ORIGEN: "XYZ"})
        resumen = build_jetsmart_resumen(build_jetsmart_guias(export), TC_JULIO_2026, VUELOS_JULIO_2026)
        self.assertEqual(set(unconfirmed_tarifa_activity(resumen)), {"XYZ"})
        self.assertEqual(resumen["estaciones"][-1]["estacion"], "XYZ")
        self.assertLess(resumen["gha_services"], 0)


@unittest.skipUnless(LIQ_FINAL.exists() and EXPORT.exists(), "faltan los archivos de tests/fixtures/jetsmart/")
class TestExportCrudo(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.export = load_jetsmart_export(EXPORT)
        cls.guias = build_jetsmart_guias(cls.export)
        cls.resumen = build_jetsmart_resumen(cls.guias, TC_JULIO_2026, VUELOS_JULIO_2026)
        cls.conciliadas = pd.read_excel(LIQ_FINAL, sheet_name="GUIAS ", usecols=range(15))
        cls.conciliadas = cls.conciliadas[cls.conciliadas["AWB"].notna()]

    def test_lee_el_export_real(self):
        self.assertEqual(len(self.export), 1961)
        period, otros_meses = detect_jetsmart_period(self.export)
        self.assertEqual(period, PERIODO)
        self.assertEqual(len(otros_meses), 61)  # guias creadas en junio 2026

    def test_prioridad_y_kg_coinciden_con_la_liquidacion_guia_por_guia(self):
        m = self.guias.merge(self.conciliadas, on="Nro guía", suffixes=("", "_ref"))
        self.assertEqual(len(m), 1925)
        pd.testing.assert_series_equal(m["GRAVADO"], m["Ingreso total flete"], check_names=False)
        pd.testing.assert_series_equal(m["Kg"], m["Kg_ref"], check_names=False)

    def test_diferencia_contra_la_liquidacion_es_solo_el_cruce_con_ariel(self):
        solo_export = self.guias[~self.guias["Nro guía"].isin(self.conciliadas["Nro guía"])]
        solo_liq = self.conciliadas[~self.conciliadas["Nro guía"].isin(self.guias["Nro guía"])]
        self.assertEqual((len(solo_export), len(solo_liq)), (36, 24))

        ref = _sheet(openpyxl.load_workbook(LIQ_FINAL, data_only=True), "LIQUIDACION")
        esperado = ref["B6"].value + solo_export["GRAVADO"].sum() - solo_liq["Ingreso total flete"].sum()
        self.assertAlmostEqual(self.resumen["ventas_netas"], esperado, places=4)


if __name__ == "__main__":
    unittest.main()
